"""Regression checks for SVG, PNG and editable native PowerPoint export."""
import io
from types import SimpleNamespace
from unittest.mock import patch

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app import app


def test_canvas_export_buttons_and_asset():
    client = app.test_client()
    response = client.get("/")
    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert 'id="download-svg"' in html
    assert 'id="download-png"' in html
    assert 'id="download-pptx"' in html
    assert 'PowerPoint (editable)' in html
    assert '/export-pptx' in html
    assert '/static/diagram_export.js' in html
    asset = client.get("/static/diagram_export.js")
    assert asset.status_code == 200
    assert b"prepareSvg" in asset.data


def test_canvas_asset_url_with_uda_prefix():
    client = app.test_client()
    response = client.get("/", headers={
        "X-Forwarded-Prefix": "/apps/mermaid-display-app",
        "X-Forwarded-Host": "example.test",
        "X-Forwarded-Proto": "https",
    })
    assert response.status_code == 200
    assert '/apps/mermaid-display-app/static/diagram_export.js' in response.get_data(as_text=True)


def test_unsaved_canvas_powerpoint_export_has_editable_shapes():
    client = app.test_client()
    response = client.post("/export-pptx", json={
        "source": "flowchart TD\n    A[Start] --> B[Done]",
        "filename": "Internal Supplier Access",
    })
    assert response.status_code == 200, response.get_data(as_text=True)[:500]
    assert response.mimetype == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert 'Internal_Supplier_Access.pptx' in response.headers["Content-Disposition"]
    deck = Presentation(io.BytesIO(response.data))
    assert len(deck.slides) == 1
    assert len(deck.slides[0].shapes) == 1
    group = deck.slides[0].shapes[0]
    # diagram-pptx deliberately groups editable shapes in one top-level object.
    assert group.shape_type == MSO_SHAPE_TYPE.GROUP
    assert len(group.shapes) >= 2
    assert any(shape.has_text_frame for shape in group.shapes)
    assert all(shape.shape_type != MSO_SHAPE_TYPE.PICTURE for shape in group.shapes)


def test_saved_repository_powerpoint_download_uses_assembled_source():
    diagram = SimpleNamespace(
        project_slug="General", name="Test Flow", revision=1,
        diagram_type="master", content="flowchart TD\n A --> B",
    )
    with (
        patch("app.load_repository_diagram", return_value=diagram),
        patch("app.renderable_repository_content", return_value="flowchart TD\n A --> B") as assembled,
    ):
        response = app.test_client().get("/repository/pptx/General/test/1")
    assert response.status_code == 200
    assembled.assert_called_once_with(diagram)
    assert 'General_Test_Flow_r1.pptx' in response.headers["Content-Disposition"]


def test_empty_or_bad_powerpoint_source_is_rejected():
    client = app.test_client()
    assert client.post("/export-pptx", json={"source": ""}).status_code == 400
    assert client.post("/export-pptx", json={"source": ["not", "text"]}).status_code == 400
    assert client.post("/export-pptx", json={"source": "x" * 300001}).status_code == 413
