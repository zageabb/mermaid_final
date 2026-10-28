"""Editable PowerPoint export for Mermaid diagrams.

Ported from mermaid-display-app. Native PowerPoint shapes and connectors, not a
rasterised screenshot, are produced by diagram-pptx where supported.
"""
import io

from diagram_pptx import render_mermaid
from pptx import Presentation
from pptx.util import Inches


def build_powerpoint(mermaid_source):
    """Render Mermaid source to a widescreen slide of editable PPTX elements."""
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    render_mermaid(mermaid_source, slide=slide, position="full")
    result = io.BytesIO()
    presentation.save(result)
    result.seek(0)
    return result
